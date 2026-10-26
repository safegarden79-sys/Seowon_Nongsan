# -*- coding: utf-8 -*-
"""마진 계산 엑셀 템플릿 만들기 → margin/template.xlsx

서버(margin/xlsx.js)가 이 템플릿을 열어 자료 시트(품목요약·일별데이터·낙찰원장)와 설정 칸을 채워 내려준다.
수식·그래프·서식은 여기서만 정한다. 계산식은 margin/margin.js 의 calc() 와 같아야 한다.

    python3 margin/build_template.py      # openpyxl 필요 (pip install openpyxl)

자료 칸의 '본보기 줄'(2행)은 서버가 서식(s 번호)을 읽어 같은 서식으로 줄을 늘리는 데 쓴다.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "template.xlsx")
SUMMARY_ROWS = 80      # 품목요약 최대 줄 (그래프 범위는 서버가 실제 줄 수로 줄인다)
CH_ROWS = (4, 11)      # 설정 채널 줄 4~11 (최대 8개)
SUM_CH = 3             # 품목요약에 마진을 펼쳐 보일 채널 수 (설정 4~6행)

GREEN = "2F7A34"
HEAD = PatternFill("solid", fgColor="E2F1D6")
INPUT = PatternFill("solid", fgColor="FFF4C2")
OUTF = PatternFill("solid", fgColor="F2F8EC")
KEY = PatternFill("solid", fgColor="FFE0C2")
thin = Side(style="thin", color="DDE7D7")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WON = '#,##0;[Red]-#,##0'
PCT = '0.0%;[Red]-0.0%'
DATE = 'yyyy-mm-dd'
KG = '0.0#'
B = Font(bold=True)
TITLE = Font(bold=True, size=15, color=GREEN)


def q(name):
    return "'" + name + "'"


S_SET, S_SUM, S_DAY, S_LOT, S_TR, S_CALC, S_CH, S_DB = "설정", "품목요약", "일별데이터", "낙찰원장", "추이", "계산기", "그래프", "품목DB"

wb = Workbook()
calc = wb.active
calc.title = S_CALC
graph = wb.create_sheet(S_CH)
summ = wb.create_sheet(S_SUM)
trend = wb.create_sheet(S_TR)
setting = wb.create_sheet(S_SET)
daily = wb.create_sheet(S_DAY)
lots = wb.create_sheet(S_LOT)
itemdb = wb.create_sheet(S_DB)


def head(ws, row, names, widths=None):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = B; c.fill = HEAD; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=row, column=i).column_letter].width = w
    ws.row_dimensions[row].height = 32


def proto(ws, row, fmts):
    """본보기 줄: 서식만 있는 빈 칸 (서버가 s 번호를 읽어 간다)"""
    for i, f in enumerate(fmts, 1):
        c = ws.cell(row=row, column=i)
        c.border = BOX
        if f: c.number_format = f


# ---------------- 설정 ----------------
st = setting
st["A1"] = "마진 계산 설정 — 판매 채널별 비용·상한 기준"; st["A1"].font = TITLE
st["A2"] = "노란 칸을 고치면 계산기·품목요약·그래프가 다시 계산됩니다. (앱 화면의 설정은 서버에 저장되고, 이 파일을 고쳐도 서버에는 반영되지 않습니다)"
head(st, 3, ["채널", "수수료율(%)", "택배비(원/박스)", "포장비(원/박스)", "기타 비용(원/박스)", "상한 기준(소매|가락)", "소매 대비 할인(%)", "가락 대비 가산(%)", "최소 마진(%)"],
     [20, 12, 14, 14, 15, 16, 14, 14, 12])
for r in range(CH_ROWS[0], CH_ROWS[1] + 1):
    for c in range(1, 10):
        cell = st.cell(row=r, column=c)
        cell.fill = INPUT; cell.border = BOX
        if c in (3, 4, 5): cell.number_format = WON
st["A13"] = "원가 기준"; st["A13"].font = B
st["C13"].fill = INPUT; st["C13"].border = BOX
st["D13"] = "최근 = 최근 경매일 가중평균 · 7일평균 = 최근 7일 가중평균 · 7일최고 = 최근 7일 중 가장 비싼 날"
st["A15"] = "자료 기간"; st["A15"].font = B
for ref in ("B15", "B16", "B17"): st[ref] = "-"   # 서버가 채운다
st["A16"] = "만든 시각"; st["A16"].font = B
st["A17"] = "시세 자료"; st["A17"].font = B
st["A19"] = "계산식"; st["A19"].font = B
for i, t in enumerate([
    "원가 C = 원가 기준에 따른 서원 낙찰 단가(박스당)",
    "부대비 F = 택배비 + 포장비 + 기타 비용,  수수료율 r",
    "손익분기가 = (C + F) ÷ (1 − r)",
    "최소 마진가 = 손익분기가 × (1 + 최소 마진%)",
    "상한가 = 소매 기준: 지역 소매 kg당 가운데값 × 박스kg × (1 − 소매 대비 할인%)  /  가락 기준: 가락시장 7일 kg당 평균 × 박스kg × (1 + 가락 대비 가산%)  (자료가 없으면 다른 기준으로 대신)",
    "최대 마진 = 상한가 × (1 − r) − C − F,   마진율 = 최대 마진 ÷ 상한가,   권장가 = (최소 마진가 + 상한가) ÷ 2",
]):
    st.cell(row=20 + i, column=1, value=t)
dv_basis = DataValidation(type="list", formula1='"최근,7일평균,7일최고"', allow_blank=False)
dv_cb = DataValidation(type="list", formula1='"소매,가락"', allow_blank=True)
st.add_data_validation(dv_basis); st.add_data_validation(dv_cb)
dv_basis.add("C13"); dv_cb.add("F4:F11")
st.freeze_panes = "B4"

SET = q(S_SET)


def set_ref(col, row):
    return f"{SET}!${col}${row}"


# ---------------- 품목요약 ----------------
sm = summ
sum_cols = ["품목", "품목군", "등급", "박스kg", "최근 경매일", "최근 단가(박스)", "7일 평균(박스)", "7일 최고(박스)", "기간 평균(박스)",
            "기간 최저 단가", "기간 최고 단가", "총 낙찰 수량", "총 낙찰 금액", "가락 7일 kg당", "전국 7일 kg당", "소매 kg당(가운데)",
            "가락 박스 환산", "가락보다 싸게 산 금액", "원가(기준)"]
for k in range(SUM_CH):
    sum_cols += [f"채널{k+1} 상한가", f"채널{k+1} 최대 마진", f"채널{k+1} 마진율"]
head(sm, 1, sum_cols, [22, 10, 7, 7, 11, 12, 12, 12, 12, 11, 11, 10, 13, 11, 11, 11, 12, 13, 12] + [12, 12, 9] * SUM_CH)
fm = [None, None, None, KG, DATE, WON, WON, WON, WON, WON, WON, '#,##0', WON, WON, WON, WON, WON, WON, WON] + [WON, WON, PCT] * SUM_CH
proto(sm, 2, fm)
sm.freeze_panes = "B2"
# 품목요약의 수식 열 (서버가 줄마다 {r} 을 바꿔 넣는다) — 템플릿 3행 이후는 비워 두고 서버가 채운다
SUM_FORMULAS = {
    "Q": "IF(N{r}>0,N{r}*D{r},\"\")",
    "R": "IF(N{r}>0,Q{r}-G{r},\"\")",
    "S": f"IF({set_ref('C',13)}=\"최근\",F{{r}},IF({set_ref('C',13)}=\"7일최고\",H{{r}},G{{r}}))",
}
cols3 = ["T", "U", "V", "W", "X", "Y", "Z", "AA", "AB"]
for k in range(SUM_CH):
    cr = CH_ROWS[0] + k
    c1, c2, c3 = cols3[k*3:k*3+3]
    retail = f"P{{r}}*D{{r}}*(1-{set_ref('G',cr)}/100)"
    garak = f"N{{r}}*D{{r}}*(1+{set_ref('H',cr)}/100)"
    SUM_FORMULAS[c1] = (f"IF({set_ref('A',cr)}=\"\",\"\",IF({set_ref('F',cr)}=\"소매\",IF(P{{r}}>0,{retail},IF(N{{r}}>0,{garak},\"\")),"
                        f"IF(N{{r}}>0,{garak},IF(P{{r}}>0,{retail},\"\"))))")
    SUM_FORMULAS[c2] = f"IF({c1}{{r}}=\"\",\"\",{c1}{{r}}*(1-{set_ref('B',cr)}/100)-S{{r}}-{set_ref('C',cr)}-{set_ref('D',cr)}-{set_ref('E',cr)})"
    SUM_FORMULAS[c3] = f"IF({c1}{{r}}=\"\",\"\",{c2}{{r}}/{c1}{{r}})"
# 채널 이름을 머리글에 보이게 (설정 A4~A6)
for k in range(SUM_CH):
    cr = CH_ROWS[0] + k
    for j, t in enumerate(["상한가", "최대 마진", "마진율"]):
        col = cols3[k*3 + j]
        sm[f"{col}1"] = f"={set_ref('A',cr)}&\" {t}\""

# ---------------- 일별데이터 ----------------
day_cols = ["날짜", "품목", "등급", "박스kg", "서원 낙찰 수량", "서원 평균 단가(박스)", "서원 최저 단가", "서원 최고 단가", "서원 낙찰 금액",
            "가락 kg당 평균", "가락 kg당 최저", "가락 kg당 최고", "가락 거래량", "전국 kg당 평균", "소매 kg당(가운데)", "가락 박스 환산", "가락보다 싸게 산 금액", "가락 시세 출처"]
head(daily, 1, day_cols, [11, 22, 7, 7, 10, 12, 11, 11, 13, 11, 11, 11, 10, 11, 11, 12, 13, 10])
proto(daily, 2, [DATE, None, None, KG, '#,##0', WON, WON, WON, WON, WON, WON, WON, '#,##0', WON, WON, WON, WON, None])
daily.freeze_panes = "C2"
daily.auto_filter.ref = "A1:R1"

# ---------------- 품목DB (서울청과·동화청과 누적 낙찰 품목) ----------------
db_cols = ["품목", "품목군", "등급", "박스kg (수량)", "처음 낙찰", "마지막 낙찰", "낙찰한 날 수", "총 수량", "총 금액", "평균 단가(박스)", "서울청과 수량", "동화청과 수량", "가락 시세 모음"]
head(itemdb, 1, db_cols, [16, 12, 8, 20, 11, 11, 10, 10, 14, 12, 11, 11, 12])
proto(itemdb, 2, [None, None, None, None, DATE, DATE, '#,##0', '#,##0', WON, WON, '#,##0', '#,##0', None])
itemdb.freeze_panes = "B2"
itemdb.auto_filter.ref = "A1:M1"

# ---------------- 낙찰원장 ----------------
lot_cols = ["경매일", "시장", "번호", "품목(사이트)", "품목(정리)", "생산자", "규격", "등급", "박스kg", "수량", "단가", "금액"]
head(lots, 1, lot_cols, [11, 10, 7, 16, 16, 14, 8, 7, 7, 7, 11, 13])
proto(lots, 2, [DATE, None, None, None, None, None, None, None, KG, '#,##0', WON, WON])
lots.freeze_panes = "A2"
lots.auto_filter.ref = "A1:L1"

# ---------------- 계산기 ----------------
c = calc
c.column_dimensions["A"].width = 3
c.column_dimensions["B"].width = 26
c.column_dimensions["C"].width = 22
c.column_dimensions["D"].width = 44
c.column_dimensions["F"].width = 14
c.column_dimensions["G"].width = 14
c.column_dimensions["H"].width = 12
c["B1"] = "서원농산 최대 마진 계산기"; c["B1"].font = TITLE
c["B2"] = "노란 칸만 고르거나 적으면 됩니다. 품목·채널은 목록에서 고릅니다."
SUMQ, DAYQ = q(S_SUM), q(S_DAY)
inputs = [("B4", "품목", None), ("B5", "판매 채널", None), ("B6", "원가 기준", None),
          ("B7", "판매가 (원/박스, 비우면 권장가)", WON), ("B8", "판매 수량 (박스)", '#,##0')]
for ref, label, f in inputs:
    c[ref] = label; c[ref].font = B
    v = c[ref.replace("B", "C")]
    v.fill = INPUT; v.border = BOX
    if f: v.number_format = f
c["C8"] = 10
c["D4"] = "← 품목요약의 품목 (등급·박스 무게별)"
c["D6"] = "최근 | 7일평균 | 7일최고"
dv_item = DataValidation(type="list", formula1="품목목록", allow_blank=False)
dv_ch = DataValidation(type="list", formula1=f"{SET}!$A$4:$A$11", allow_blank=False)
dv_b2 = DataValidation(type="list", formula1='"최근,7일평균,7일최고"', allow_blank=False)
for dv, ref in ((dv_item, "C4"), (dv_ch, "C5"), (dv_b2, "C6")):
    c.add_data_validation(dv); dv.add(ref)
wb.defined_names["품목목록"] = DefinedName("품목목록", attr_text=f"OFFSET({SUMQ}!$A$2,0,0,MAX(1,COUNTA({SUMQ}!$A:$A)-1),1)")

# 도움 칸 (H 열): 품목 줄 번호·채널 줄 번호·상한 후보
c["F3"] = "계산용 (손대지 않음)"; c["F3"].font = Font(italic=True, color="7A857D")
helpers = [
    ("F4", "품목 줄", "H4", f"IFERROR(MATCH($C$4,{SUMQ}!$A:$A,0),0)"),
    ("F5", "채널 줄", "H5", f"IFERROR(MATCH($C$5,{SET}!$A$4:$A$11,0),0)"),
    ("F6", "소매 상한", "H6", f"IF(AND($H$4>0,$H$5>0,$C$15>0),$C$15*$C$11*(1-INDEX({SET}!$G$4:$G$11,$H$5)/100),\"\")"),
    ("F7", "가락 상한", "H7", f"IF(AND($H$4>0,$H$5>0,$C$14>0),$C$14*$C$11*(1+INDEX({SET}!$H$4:$H$11,$H$5)/100),\"\")"),
    ("F8", "상한 기준", "H8", f"IF($H$5>0,INDEX({SET}!$F$4:$F$11,$H$5),\"\")"),
]
for lref, lt, vref, f in helpers:
    c[lref] = lt; c[lref].font = Font(color="7A857D")
    c[vref] = "=" + f; c[vref].font = Font(color="7A857D")
c["H6"].number_format = WON; c["H7"].number_format = WON


def sidx(col):
    return f"INDEX({SUMQ}!${col}:${col},$H$4)"


def cidx(col):
    return f"INDEX({SET}!${col}$4:${col}$11,$H$5)"


rows = [
    (10, "품목 정보", None, None, None),
    (11, "박스 무게(kg)", f"IF($H$4>0,{sidx('D')},\"\")", KG, ""),
    (12, "최근 경매일", f"IF($H$4>0,{sidx('E')},\"\")", DATE, ""),
    (13, "원가 (박스당 낙찰 단가)", f"IF($H$4=0,\"\",IF($C$6=\"최근\",{sidx('F')},IF($C$6=\"7일최고\",{sidx('H')},{sidx('G')})))", WON, "원가 기준에 따라 최근·7일평균·7일최고"),
    (14, "가락시장 7일 kg당 평균", f"IF($H$4>0,N({sidx('N')}),0)", WON, "공영도매시장 경락가 중 가락시장 (특·상만)"),
    (15, "지역 소매 kg당 가운데값", f"IF($H$4>0,N({sidx('P')}),0)", WON, "KAMIS 지역별 소매가"),
    (17, "판매 채널 비용", None, None, None),
    (18, "수수료율", f"IF($H$5>0,{cidx('B')}/100,0)", PCT, ""),
    (19, "부대비 (택배+포장+기타)", f"IF($H$5>0,{cidx('C')}+{cidx('D')}+{cidx('E')},0)", WON, ""),
    (21, "결과", None, None, None),
    (22, "손익분기가", "IF($C$13=\"\",\"\",($C$13+$C$19)/(1-$C$18))", WON, "이 값보다 싸게 팔면 손해"),
    (23, "최소 마진가 (하한)", f"IF($C$22=\"\",\"\",$C$22*(1+IF($H$5>0,{cidx('I')},0)/100))", WON, "최소 마진을 지키는 가장 낮은 값"),
    (24, "상한가", "IF($C$13=\"\",\"\",IF($H$8=\"소매\",IF($H$6<>\"\",$H$6,$H$7),IF($H$7<>\"\",$H$7,$H$6)))", WON, "이보다 비싸면 소비자가 마트·시장을 택할 값"),
    (25, "상한 기준", "IF($C$24=\"\",\"자료 없음\",IF($H$8=\"소매\",IF($H$6<>\"\",\"소매\",\"가락(대신)\"),IF($H$7<>\"\",\"가락\",\"소매(대신)\")))", None, ""),
    (26, "최대 마진 (상한가에 팔 때, 박스당)", "IF(OR($C$24=\"\",$C$13=\"\"),\"\",$C$24*(1-$C$18)-$C$13-$C$19)", WON, "이 품목·채널에서 낼 수 있는 가장 큰 마진"),
    (27, "최대 마진율", "IF($C$26=\"\",\"\",$C$26/$C$24)", PCT, ""),
    (28, "권장가", "IF($C$23=\"\",\"\",IF($C$24=\"\",$C$23,IF($C$24>$C$23,($C$23+$C$24)/2,$C$23)))", WON, "하한과 상한의 가운데"),
    (29, "권장가 마진 (박스당)", "IF($C$28=\"\",\"\",$C$28*(1-$C$18)-$C$13-$C$19)", WON, ""),
    (31, "내 판매가로 계산", None, None, None),
    (32, "적용 판매가", "IF(N($C$7)>0,$C$7,$C$28)", WON, "판매가를 비우면 권장가"),
    (33, "박스당 마진", "IF(OR($C$32=\"\",$C$13=\"\"),\"\",$C$32*(1-$C$18)-$C$13-$C$19)", WON, ""),
    (34, "마진율", "IF($C$33=\"\",\"\",$C$33/$C$32)", PCT, ""),
    (35, "총 마진 (× 판매 수량)", "IF($C$33=\"\",\"\",$C$33*N($C$8))", WON, ""),
    (36, "최대 마진과의 차이 (박스당)", "IF(OR($C$33=\"\",$C$26=\"\"),\"\",$C$33-$C$26)", WON, "0 에 가까울수록 최대 마진에 가깝게 판 것"),
]
for r, label, f, fmt, note in rows:
    cell = c.cell(row=r, column=2, value=label)
    if f is None:
        cell.font = Font(bold=True, color=GREEN, size=12)
        continue
    v = c.cell(row=r, column=3, value="=" + f)
    v.fill = OUTF; v.border = BOX
    if fmt: v.number_format = fmt
    if note: c.cell(row=r, column=4, value=note).font = Font(color="7A857D")
for r in (22, 24, 26, 28, 33, 35):
    c.cell(row=r, column=3).font = B
c["C26"].fill = KEY

# 가격별 마진 표 (그래프용) — 손익분기가의 90% ~ 상한가의 110%
c["B39"] = "가격별 마진 (그래프 시트의 '가격별 마진' 그래프)"; c["B39"].font = Font(bold=True, color=GREEN, size=12)
for i, t in enumerate(["판매가", "박스당 마진", "마진율"]):
    h = c.cell(row=40, column=2 + i, value=t); h.font = B; h.fill = HEAD; h.border = BOX
STEPS = 20
lo = "($C$22*0.9)"
hi = "(IF($C$24=\"\",$C$22*1.6,MAX($C$24,$C$22)*1.1))"
for i in range(STEPS + 1):
    r = 41 + i
    p = c.cell(row=r, column=2, value=f"=IF($C$22=\"\",NA(),ROUND({lo}+({hi}-{lo})*{i}/{STEPS},-1))")
    m = c.cell(row=r, column=3, value=f"=IF($C$22=\"\",NA(),B{r}*(1-$C$18)-$C$13-$C$19)")
    rt = c.cell(row=r, column=4, value=f"=IF($C$22=\"\",NA(),C{r}/B{r})")
    p.number_format = WON; m.number_format = WON; rt.number_format = PCT
    for x in (p, m, rt): x.border = BOX
c.freeze_panes = "A4"

# ---------------- 추이 (선택한 품목, 최근 40일) ----------------
t = trend
t["A1"] = "선택한 품목의 날짜별 추이 (계산기 C4 의 품목)"; t["A1"].font = TITLE
t["F1"] = "품목"; t["G1"] = f"={q(S_CALC)}!$C$4"; t["G1"].font = B
head(t, 2, ["날짜", "서원 낙찰 평균(박스)", "가락 평균(박스 환산)", "전국 평균(박스 환산)", "소매 기준 상한(박스)", "가락보다 싸게 산 금액", "그날 낙찰 수량"],
     [11, 16, 16, 16, 16, 16, 12])
TREND_DAYS = 40
CK = f"{q(S_CALC)}!$C$4"
for i in range(TREND_DAYS):
    r = 3 + i
    t.cell(row=r, column=1, value=f"=IFERROR(MAX({DAYQ}!$A:$A),TODAY())-{TREND_DAYS-1-i}").number_format = DATE
    crit = f"{DAYQ}!$B:$B,$G$1,{DAYQ}!$A:$A,$A{r}"
    t.cell(row=r, column=2, value=f"=IFERROR(SUMIFS({DAYQ}!$I:$I,{crit})/SUMIFS({DAYQ}!$E:$E,{crit}),NA())")
    t.cell(row=r, column=3, value=f"=IFERROR(AVERAGEIFS({DAYQ}!$J:$J,{crit})*{q(S_CALC)}!$C$11,NA())")
    t.cell(row=r, column=4, value=f"=IFERROR(AVERAGEIFS({DAYQ}!$N:$N,{crit})*{q(S_CALC)}!$C$11,NA())")
    t.cell(row=r, column=5, value=f"=IFERROR(AVERAGEIFS({DAYQ}!$O:$O,{crit})*{q(S_CALC)}!$C$11*(1-IF({q(S_CALC)}!$H$5>0,INDEX({SET}!$G$4:$G$11,{q(S_CALC)}!$H$5),0)/100),NA())")
    t.cell(row=r, column=6, value=f"=IFERROR(C{r}-B{r},NA())")
    t.cell(row=r, column=7, value=f"=SUMIFS({DAYQ}!$E:$E,{crit})")
    for col in range(1, 8):
        t.cell(row=r, column=col).border = BOX
        if col > 1: t.cell(row=r, column=col).number_format = WON if col < 7 else '#,##0'
t["A44"] = "빈 칸(#N/A)은 그날 자료가 없다는 뜻입니다 — 그래프에서는 건너뜁니다."
t.freeze_panes = "B3"

# ---------------- 그래프 ----------------
g = graph
g["A1"] = "그래프 — 계산기에서 품목·채널을 바꾸면 함께 바뀝니다"; g["A1"].font = TITLE

ch1 = LineChart()
ch1.title = "선택 품목: 서원 낙찰 단가 vs 가락·전국 시세 (박스당)"
ch1.y_axis.title = "원/박스"; ch1.y_axis.numFmt = '#,##0'
ch1.height, ch1.width = 9, 24
data = Reference(trend, min_col=2, max_col=5, min_row=2, max_row=2 + TREND_DAYS)
ch1.add_data(data, titles_from_data=True)
ch1.set_categories(Reference(trend, min_col=1, min_row=3, max_row=2 + TREND_DAYS))
ch1.x_axis.numFmt = "mm/dd"
for s in ch1.series:
    s.smooth = False
    s.marker.symbol = "circle"; s.marker.size = 4
ch1.x_axis.delete = False; ch1.y_axis.delete = False
g.add_chart(ch1, "A3")

ch2 = BarChart()
ch2.title = "선택 품목: 가락 평균보다 싸게 산 금액 (박스당, + 이면 싸게 산 날)"
ch2.y_axis.numFmt = '#,##0'; ch2.height, ch2.width = 7.5, 24
ch2.add_data(Reference(trend, min_col=6, min_row=2, max_row=2 + TREND_DAYS), titles_from_data=True)
ch2.set_categories(Reference(trend, min_col=1, min_row=3, max_row=2 + TREND_DAYS))
ch2.x_axis.numFmt = "mm/dd"; ch2.legend = None
ch2.x_axis.delete = False; ch2.y_axis.delete = False
g.add_chart(ch2, "A22")

ch3 = LineChart()
ch3.title = "가격별 마진 (선택 품목·채널) — 0 위로 올라가는 지점이 손익분기"
ch3.y_axis.numFmt = '#,##0'; ch3.height, ch3.width = 7.5, 24
ch3.add_data(Reference(calc, min_col=3, min_row=40, max_row=41 + STEPS), titles_from_data=True)
ch3.set_categories(Reference(calc, min_col=2, min_row=41, max_row=41 + STEPS))
ch3.x_axis.numFmt = '#,##0'; ch3.legend = None
ch3.x_axis.delete = False; ch3.y_axis.delete = False
g.add_chart(ch3, "A38")

ch4 = BarChart()
ch4.type = "bar"
ch4.title = "품목별 최대 마진 (박스당) — 채널 1~3"
ch4.y_axis.numFmt = '#,##0'; ch4.height, ch4.width = 16, 24
for k in range(SUM_CH):
    col = 21 + k * 3   # U, X, AA
    ch4.add_data(Reference(summ, min_col=col, min_row=1, max_row=1 + SUMMARY_ROWS), titles_from_data=True)
ch4.set_categories(Reference(summ, min_col=1, min_row=2, max_row=1 + SUMMARY_ROWS))
ch4.x_axis.delete = False; ch4.y_axis.delete = False
g.add_chart(ch4, "P3")

ch5 = BarChart()
ch5.type = "bar"
ch5.title = "품목별 서원 7일 평균 단가 vs 가락 박스 환산"
ch5.y_axis.numFmt = "#,##0"; ch5.height, ch5.width = 16, 24
ch5.add_data(Reference(summ, min_col=7, min_row=1, max_row=1 + SUMMARY_ROWS), titles_from_data=True)
ch5.add_data(Reference(summ, min_col=17, min_row=1, max_row=1 + SUMMARY_ROWS), titles_from_data=True)
ch5.set_categories(Reference(summ, min_col=1, min_row=2, max_row=1 + SUMMARY_ROWS))
ch5.x_axis.delete = False; ch5.y_axis.delete = False
g.add_chart(ch5, "P36")

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
calc.sheet_properties.tabColor = "2F7A34"
graph.sheet_properties.tabColor = "D9472B"
setting.sheet_properties.tabColor = "FFC943"

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)

# 서버가 쓰는 수식 열 정의를 JSON 으로도 남긴다 (xlsx.js 가 읽음)
import json
with open(os.path.join(os.path.dirname(OUT), "template.json"), "w", encoding="utf-8") as fp:
    json.dump({"summaryRows": SUMMARY_ROWS, "summaryFormulas": SUM_FORMULAS, "sumChannels": SUM_CH, "trendDays": TREND_DAYS,
               "sheets": {"calc": S_CALC, "graph": S_CH, "summary": S_SUM, "trend": S_TR, "settings": S_SET, "daily": S_DAY, "lots": S_LOT, "itemdb": S_DB}},
              fp, ensure_ascii=False, indent=1)
print("만듦:", OUT)
